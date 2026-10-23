"""问答附件的私有存储与文本提取。

附件只作为用户当前会话的提示词上下文，不进入公共向量库。原文件与提取文本
按用户隔离存放，前端只持有不可猜测的 UUID。
"""

from __future__ import annotations

import base64
import csv
import hashlib
import json
import re
import shutil
import uuid
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from ai_engine.config import settings
from ai_engine.rag.loaders import UnsupportedDocument, load_document

MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_FILES_PER_UPLOAD = 8
MAX_EXTRACTED_CHARS = 50_000
# 图片随 deepseek-flash 的同一次对话请求内联。单文件沿用 15 MB（低于接口 32 MB），
# 多张合计再卡 30 MB，避免 base64 后顶满 48 MB 的请求体。
MAX_REQUEST_IMAGE_BYTES = 30 * 1024 * 1024
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".gif", ".webp"}
_IMAGE_MIME = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".gif": "image/gif",
    ".webp": "image/webp",
}
SUPPORTED_SUFFIXES = {
    ".pdf",
    ".docx",
    ".doc",
    ".pptx",
    ".xlsx",
    ".csv",
    ".tsv",
    ".md",
    ".txt",
    ".html",
    ".htm",
    *IMAGE_SUFFIXES,
}
_ATTACHMENT_ID = re.compile(r"^[0-9a-f]{32}$")


def _user_key(user_id: str) -> str:
    return hashlib.sha256(user_id.encode("utf-8")).hexdigest()[:24]


def _user_root(user_id: str) -> Path:
    root = settings.upload_dir / "qa_attachments" / _user_key(user_id)
    root.mkdir(parents=True, exist_ok=True)
    return root


def _attachment_dir(user_id: str, attachment_id: str) -> Path | None:
    if not _ATTACHMENT_ID.fullmatch(attachment_id or ""):
        return None
    return _user_root(user_id) / attachment_id


def _detect_image_format(content: bytes) -> str | None:
    """按文件头判断真实格式。DeepSeek 也是看内容，不看扩展名或 MIME。"""
    if content.startswith(b"\xff\xd8\xff"):
        return "jpeg"
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if content.startswith((b"GIF87a", b"GIF89a")):
        return "gif"
    if len(content) >= 12 and content[:4] == b"RIFF" and content[8:12] == b"WEBP":
        return "webp"
    return None


def _expected_image_format(suffix: str) -> str:
    if suffix in {".jpg", ".jpeg"}:
        return "jpeg"
    return suffix.lstrip(".")


def _load_pptx(path: Path) -> str:
    from pptx import Presentation

    parts: list[str] = []
    presentation = Presentation(str(path))
    for index, slide in enumerate(presentation.slides, start=1):
        slide_parts: list[str] = []
        for shape in slide.shapes:
            text = (getattr(shape, "text", "") or "").strip()
            if text:
                slide_parts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        slide_parts.append(" | ".join(cells))
        if slide_parts:
            parts.append(f"[第 {index} 页]\n" + "\n".join(slide_parts))
    return "\n\n".join(parts)


def _load_delimited(path: Path, delimiter: str) -> str:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        lines: list[str] = []
        for row_index, row in enumerate(reader, start=1):
            if row_index > 1000:
                lines.append("[仅提取前 1000 行]")
                break
            lines.append(f"{row_index}: " + " | ".join(cell.strip() for cell in row))
    return "\n".join(lines)


def _load_xlsx(path: Path) -> str:
    """用标准库读取 xlsx 的单元格值，避免把公式或格式当作计算结果。"""
    main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    office_rel_ns = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > 10_000 or sum(item.file_size for item in members) > 100_000_000:
            raise UnsupportedDocument("表格解压后的内容过大，无法安全解析")

        shared: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in root.findall(f"{{{main_ns}}}si"):
                shared.append("".join(node.text or "" for node in item.iter(f"{{{main_ns}}}t")))

        workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        relationships = ElementTree.fromstring(
            archive.read("xl/_rels/workbook.xml.rels")
        )
        targets = {
            rel.attrib.get("Id", ""): rel.attrib.get("Target", "")
            for rel in relationships.findall(f"{{{rel_ns}}}Relationship")
        }
        parts: list[str] = []
        for sheet in workbook.findall(f".//{{{main_ns}}}sheet"):
            name = sheet.attrib.get("name", "未命名工作表")
            rel_id = sheet.attrib.get(f"{{{office_rel_ns}}}id", "")
            target = targets.get(rel_id, "")
            member = target.lstrip("/")
            if not member.startswith("xl/"):
                member = f"xl/{member}"
            if member not in archive.namelist():
                continue
            root = ElementTree.fromstring(archive.read(member))
            lines: list[str] = []
            for row_index, row in enumerate(root.findall(f".//{{{main_ns}}}row"), start=1):
                if row_index > 1000:
                    lines.append("[仅提取前 1000 行]")
                    break
                values: list[str] = []
                for cell in row.findall(f"{{{main_ns}}}c"):
                    kind = cell.attrib.get("t", "")
                    coordinate = cell.attrib.get("r", "")
                    formula = cell.findtext(f"{{{main_ns}}}f")
                    if kind == "inlineStr":
                        value = "".join(
                            node.text or ""
                            for node in cell.iter(f"{{{main_ns}}}t")
                        )
                    else:
                        value = cell.findtext(f"{{{main_ns}}}v") or ""
                        if kind == "s" and value.isdigit():
                            index = int(value)
                            value = shared[index] if index < len(shared) else value
                        elif kind == "b":
                            value = "TRUE" if value == "1" else "FALSE"
                    if formula:
                        value = f"{value} [公式={formula}]".strip()
                    values.append(f"{coordinate}={value}")
                if values:
                    lines.append(" | ".join(values))
            parts.append(f"[工作表：{name}]\n" + "\n".join(lines))
    return "\n\n".join(parts)


def _validate_office_archive(path: Path) -> None:
    """拒绝异常膨胀的 Office ZIP，避免小文件在解析时占满内存。"""
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > 10_000 or sum(item.file_size for item in members) > 100_000_000:
                raise UnsupportedDocument("Office 文件解压后的内容过大，无法安全解析")
    except zipfile.BadZipFile as exc:
        raise UnsupportedDocument("Office 文件已损坏或格式不正确") from exc


def extract_text(path: Path) -> tuple[str, list[str]]:
    suffix = path.suffix.lower()
    warnings: list[str] = []
    if suffix in {".docx", ".pptx", ".xlsx"}:
        _validate_office_archive(path)
    if suffix in {".csv", ".tsv"}:
        text = _load_delimited(path, "," if suffix == ".csv" else "\t")
    elif suffix == ".xlsx":
        text = _load_xlsx(path)
    elif suffix == ".pptx":
        text = _load_pptx(path)
        warnings.append("演示文稿仅提取页面文字和表格，不解析图片、动画与版式关系。")
    else:
        text = load_document(path)

    text = (text or "").strip()
    if not text:
        raise UnsupportedDocument(
            "没有提取到可读文字。扫描版 PDF 和手写内容需要先做 OCR；"
            "照片请直接上传 JPG、PNG、GIF 或 WebP。"
        )
    if len(text) > MAX_EXTRACTED_CHARS:
        text = text[:MAX_EXTRACTED_CHARS]
        warnings.append(
            f"文件内容较长，本次保留前 {MAX_EXTRACTED_CHARS:,} 个字符用于问答。"
        )
    return text, warnings


def save_attachment(
    *, user_id: str, original_name: str, content: bytes, content_type: str = ""
) -> dict:
    suffix = Path(original_name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        allowed = "、".join(sorted(SUPPORTED_SUFFIXES))
        raise UnsupportedDocument(f"不支持 {suffix or '无扩展名'} 文件；目前支持：{allowed}")
    if not content:
        raise UnsupportedDocument("文件为空")
    if len(content) > MAX_FILE_BYTES:
        raise UnsupportedDocument("单个文件不能超过 15 MB")
    image_format = ""
    stored_type = content_type
    if suffix in IMAGE_SUFFIXES:
        image_format = _detect_image_format(content) or ""
        if image_format != _expected_image_format(suffix):
            raise UnsupportedDocument(
                "图片内容与扩展名不一致，或不是有效的 JPEG、PNG、GIF、WebP"
            )
        stored_type = _IMAGE_MIME[suffix]

    attachment_id = uuid.uuid4().hex
    folder = _user_root(user_id) / attachment_id
    folder.mkdir(parents=True, exist_ok=False)
    source_path = folder / f"original{suffix}"
    try:
        source_path.write_bytes(content)
        if image_format:
            text, warnings = "", []
        else:
            text, warnings = extract_text(source_path)
            (folder / "extracted.txt").write_text(text, encoding="utf-8")
        metadata = {
            "id": attachment_id,
            "name": Path(original_name).name[:512] or f"未命名{suffix}",
            "size": len(content),
            "content_type": stored_type,
            "suffix": suffix,
            "kind": "image" if image_format else "text",
            "status": "ready",
            "char_count": len(text),
            "preview": text[:240],
            "warnings": warnings,
        }
        (folder / "metadata.json").write_text(
            json.dumps(metadata, ensure_ascii=False), encoding="utf-8"
        )
        return metadata
    except Exception:
        shutil.rmtree(folder, ignore_errors=True)
        raise


def original_path(user_id: str, attachment_id: str) -> Path | None:
    folder = _attachment_dir(user_id, attachment_id)
    if not folder or not folder.is_dir():
        return None
    matches = sorted(folder.glob("original.*"))
    return matches[0] if matches else None


def image_data_url(path: Path, mime: str) -> str:
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{mime};base64,{encoded}"


def load_attachment(user_id: str, attachment_id: str) -> tuple[dict, str] | None:
    folder = _attachment_dir(user_id, attachment_id)
    if not folder or not folder.is_dir():
        return None
    try:
        metadata = json.loads((folder / "metadata.json").read_text(encoding="utf-8"))
        if metadata.get("kind") == "image":
            return metadata, ""
        text = (folder / "extracted.txt").read_text(encoding="utf-8")
    except (OSError, json.JSONDecodeError):
        return None
    return metadata, text


def delete_attachment(user_id: str, attachment_id: str) -> bool:
    folder = _attachment_dir(user_id, attachment_id)
    if not folder or not folder.is_dir():
        return False
    shutil.rmtree(folder)
    return True
