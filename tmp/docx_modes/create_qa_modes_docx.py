from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/huangkefan/Desktop/实验室/北大问教2.0")
OUTPUT = ROOT / "docs" / "北大问教四种问答模式设计说明.docx"

FONT_CN = "PingFang SC"
FONT_LATIN = "Aptos"
NAVY = "1F4E78"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F6F7"
BORDER = "D9D9D9"
BLACK = "000000"
MUTED = "606266"


MODES = [
    {
        "name": "通用问答",
        "code": "ask",
        "summary": "解释教学问题，并给出有依据的可执行建议",
        "audience": (
            "面向有一般教学咨询需求的高校教师。适合解释教学概念、查找教学策略和"
            "回答局部问题，例如“大班课如何提高参与度”和“教学目标怎样写得可测量”。"
        ),
        "tasks": [
            "教学概念解释",
            "通用教学策略建议",
            "教学难题咨询",
            "理论与实践案例查询",
        ],
        "workflow": [
            "根据最近的对话补全当前问题中的指代。",
            "生成 2–4 条语义检索语句，覆盖核心教学概念、适用情境和实践策略。",
            "在默认知识库中执行分块向量检索。",
            "按相似度过滤、合并和去重检索文本块。",
            "由大模型结合检索证据回答，并给出可执行建议。",
        ],
        "knowledge": [
            "教育学基础理论",
            "北大教学沙龙",
            "课程大纲参考",
            "学科教学知识",
            "国内外高校教学案例",
            "Stanford AIMES",
        ],
        "result": [
            "直接结论",
            "理论或案例依据",
            "可执行建议",
            "适用条件或实施风险",
        ],
        "format": (
            "Markdown 自然结构，不强制固定模板。引用资料时使用 [1]、[2] 等编号；"
            "页面下方展示来源、相似度和原文网址。"
        ),
        "status": "当前使用常规分块 RAG。",
    },
    {
        "name": "大纲设计",
        "code": "syllabus_design",
        "summary": "匹配内容相近课程，设计目标、内容、进度与评价体系",
        "audience": (
            "面向准备新开课、修订课程大纲或重新设计课程体系的教师、院系负责人和课程团队。"
            "适合匹配内容相近课程，并设计课程目标、周计划、教学形式及考核结构。"
        ),
        "tasks": [
            "新课程大纲设计",
            "现有大纲修订",
            "相似国外课程匹配",
            "目标、内容、活动和评价一致性设计",
        ],
        "workflow": [
            "提取课程主题、核心概念、学生层次、先修知识和课程类型。",
            "生成中英文课程匹配查询，补充学习目标、课程组织和考核方式等维度。",
            "优先检索中文课程检索画像和课程资料，确定最多 3 门相似 MIT 课程。",
            "按课程目录读取完整相关文档，不仅使用命中的短文本块。",
            "补充检索课程设计理论和学科教学知识。",
            "综合形成大纲方案，并区分“来源事实”和“本次设计建议”。",
        ],
        "knowledge": [
            "国内外高校教学案例，主要用于匹配 MIT OCW 课程",
            "课程大纲参考",
            "学科教学知识",
            "教育学基础理论",
        ],
        "result": [
            "需求理解与必要假设",
            "相似课程及匹配理由",
            "可迁移的设计要素",
            "课程定位与学生对象",
            "可测量的课程目标",
            "核心内容与周次安排",
            "教学组织形式",
            "作业与考核结构",
            "目标、活动和评价一致性表",
            "原始课程网址",
        ],
        "format": (
            "Markdown 结构化输出。界面单独展示匹配课程、原文网址、匹配度和实际读取的课程文档。"
        ),
        "status": (
            "当前对 MIT OCW 实施两阶段检索：先匹配课程，再加载完整相关文档。"
            "每门课程按文档边界读取，不在文档中间截断。"
        ),
    },
    {
        "name": "教学设计",
        "code": "teaching_design",
        "summary": "面向一节课或一个单元设计活动、互动与形成性评价",
        "audience": (
            "面向需要设计一节课、一个知识单元或一项教学活动的教师。适合 90 分钟课堂设计、"
            "难点讲解、课堂互动、形成性评价和 AI 作业设计。"
        ),
        "tasks": [
            "单课时或教学单元设计",
            "抽象概念和典型难点讲解",
            "课堂互动与学生活动设计",
            "形成性评价设计",
            "生成式 AI 教学和作业设计",
        ],
        "workflow": [
            "识别教学主题、学生情况、课时长度和教学难点。",
            "生成覆盖知识内容、学习难点、教学活动和评价方法的检索语句。",
            "匹配知识主题或教学方法相近的课程案例。",
            "对命中的 MIT 课程加载完整相关材料。",
            "补充检索北大教学案例、教学沙龙、学科教学方法和教育理论。",
            "如问题涉及生成式 AI，补充检索 AIMES 的 AI 教学策略和课程政策。",
            "将案例经验转换为适合当前课堂的教学流程。",
        ],
        "knowledge": [
            "国内外高校教学案例",
            "北大教学沙龙",
            "学科教学知识",
            "教育学基础理论",
            "Stanford AIMES",
        ],
        "result": [
            "学习目标",
            "学情假设与主要难点",
            "教学设计原则",
            "分阶段教学流程表",
            "每阶段的时间、教师行为、学生活动和预期产出",
            "课堂提问与互动设计",
            "形成性评价方法",
            "课后任务",
            "实施风险与备选方案",
            "参考案例及原始网址",
        ],
        "format": (
            "Markdown 结构化输出。课时信息充分时，优先使用时间轴或教学流程表；"
            "引用案例时附来源编号和原文网址。"
        ),
        "status": (
            "当前对 MIT OCW 课程执行案例匹配和完整文档加载；"
            "其他教学案例仍使用分块 RAG。"
        ),
    },
    {
        "name": "学情诊断",
        "code": "learning_diagnosis",
        "summary": "根据作业、课堂表现或评估报告识别问题并提出干预方案",
        "audience": (
            "面向已经掌握学生作业、测验、课堂表现或课堂评估报告，希望发现问题并制定干预方案的教师。"
            "适合错误分析、学生分层、薄弱知识点定位、课堂问题诊断和改进计划。"
        ),
        "tasks": [
            "作业或测验错误分析",
            "薄弱知识点与学习困难定位",
            "学生分层与针对性干预",
            "历史课堂评估报告解读",
            "教学改进计划设计",
        ],
        "workflow": [
            "优先读取教师提供的学生表现数据或指定的历史课堂评估报告。",
            "界定数据范围、样本限制以及根据当前证据不能得出的结论。",
            "提取可直接观察的现象，如错误类型、得分分布、完成情况和课堂表现。",
            "将事实证据和原因推断分开，并标明诊断置信程度。",
            "检索学习科学、错误诊断和学科教学知识。",
            "形成问题分类、原因假设和学生分组。",
            "给出分层干预方案及下一轮验证方法。",
        ],
        "knowledge": [
            "教师指定的历史评估报告，作为首要证据",
            "教育学基础理论",
            "学科教学知识",
            "北大教学沙龙",
        ],
        "result": [
            "数据范围与诊断限制",
            "关键证据",
            "主要问题及严重程度",
            "可观察现象与原因假设对照表",
            "诊断置信程度",
            "学生分组或问题分类",
            "针对性教学干预",
            "下一轮测评或验证方案",
        ],
        "format": (
            "Markdown 结构化诊断报告。没有学生数据时，不直接下诊断结论，"
            "只返回诊断框架和所需材料清单。"
        ),
        "status": (
            "当前支持粘贴或概述学生表现，以及选择历史课堂评估报告。"
            "作业文件批量上传、表格统计和学生匿名化处理尚未接入。"
        ),
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=BORDER, size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Inches(inches))))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size: float | None = None, bold: bool | None = None, color=BLACK) -> None:
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, *, space_after=4, line_spacing=1.15) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing


def set_cell_text(cell, text: str, *, bold=False, color=BLACK, size=10.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    style_paragraph(paragraph, space_after=0, line_spacing=1.12)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_borders(cell)


def add_cell_list(cell, items: list[str], *, numbered=False, size=10.3) -> None:
    cell.text = ""
    for index, item in enumerate(items, 1):
        paragraph = cell.paragraphs[0] if index == 1 else cell.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.03)
        paragraph.paragraph_format.first_line_indent = Inches(-0.01)
        style_paragraph(paragraph, space_after=2.5, line_spacing=1.12)
        prefix = f"{index}. " if numbered else "• "
        run = paragraph.add_run(prefix + item)
        set_run_font(run, size=size, color=BLACK)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_borders(cell)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    title = styles["Title"]
    title.font.name = FONT_LATIN
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 16, 16, 8),
        ("Heading 2", 13, 12, 6),
    ):
        style = styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_summary_table(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (1.15, 2.35, 3.55)
    headers = ("模式", "面向对象", "核心工作流")
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF", size=10.5)
        set_cell_shading(table.rows[0].cells[i], NAVY)
        set_cell_width(table.rows[0].cells[i], widths[i])
        table.rows[0].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_repeat_table_header(table.rows[0])

    summaries = [
        ("通用问答", "有一般教学咨询需求的高校教师", "生成检索语句 → 分块 RAG → 带引用的可执行建议"),
        ("大纲设计", "新开课、修订大纲的教师和课程团队", "匹配相似课程 → 读取完整课程文档 → 设计大纲"),
        ("教学设计", "设计单课时、教学单元或活动的教师", "匹配教学案例 → 读取完整材料 → 生成课时流程"),
        ("学情诊断", "持有作业、测验、课堂表现或评估报告的教师", "证据优先 → 现象与推断分离 → 分层干预和验证"),
    ]
    for row_index, row_data in enumerate(summaries, 1):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_data):
            set_cell_text(row.cells[i], value, bold=(i == 0), size=10.2)
            set_cell_width(row.cells[i], widths[i])
            if row_index % 2 == 0:
                set_cell_shading(row.cells[i], PALE_BLUE)
    doc.add_paragraph()


def add_mode_table(doc: Document, mode: dict) -> None:
    heading = doc.add_heading(mode["name"], level=1)
    heading.paragraph_format.page_break_before = True

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.keep_with_next = True
    style_paragraph(subtitle, space_after=8, line_spacing=1.15)
    run = subtitle.add_run(f"{mode['code']}  {mode['summary']}")
    set_run_font(run, size=10.5, color=MUTED)

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (1.25, 5.8)
    for i, value in enumerate(("设计项", "说明")):
        set_cell_text(table.rows[0].cells[i], value, bold=True, color="FFFFFF", size=10.5)
        set_cell_shading(table.rows[0].cells[i], NAVY)
        set_cell_width(table.rows[0].cells[i], widths[i])
        table.rows[0].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_repeat_table_header(table.rows[0])

    rows = [
        ("面向对象", mode["audience"], "text"),
        ("典型任务", mode["tasks"], "bullets"),
        ("工作流", mode["workflow"], "numbered"),
        ("默认知识库", mode["knowledge"], "bullets"),
        ("返回内容", mode["result"], "numbered"),
        ("返回格式", mode["format"], "text"),
        ("当前边界", mode["status"], "text"),
    ]
    for row_index, (label, value, kind) in enumerate(rows, 1):
        row = table.add_row()
        prevent_row_split(row)
        set_cell_text(row.cells[0], label, bold=True, size=10.3)
        set_cell_shading(row.cells[0], PALE_GRAY)
        if kind == "text":
            set_cell_text(row.cells[1], value, size=10.3)
        else:
            add_cell_list(row.cells[1], value, numbered=(kind == "numbered"), size=10.1)
        set_cell_width(row.cells[0], widths[0])
        set_cell_width(row.cells[1], widths[1])
        if row_index % 2 == 0:
            set_cell_shading(row.cells[1], PALE_BLUE)


def add_loading_table(doc: Document) -> None:
    doc.add_heading("完整课程材料装载", level=1)
    p = doc.add_paragraph(
        "大纲设计和教学设计都使用两阶段检索。系统先根据检索摘要或课程文本确定候选课程，"
        "再从课程目录读取完整的相关文档。当前每次最多装载 3 门 MIT 课程。"
    )
    style_paragraph(p, space_after=8, line_spacing=1.25)

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (1.45, 5.6)
    for i, value in enumerate(("模式", "按优先级装载的文档")):
        set_cell_text(table.rows[0].cells[i], value, bold=True, color="FFFFFF", size=10.5)
        set_cell_shading(table.rows[0].cells[i], NAVY)
        set_cell_width(table.rows[0].cells[i], widths[i])
        table.rows[0].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_repeat_table_header(table.rows[0])
    data = [
        (
            "大纲设计",
            "检索摘要.md、course-description.md、syllabus.md、calendar.md、assignments.md、instructor-insights.md",
        ),
        (
            "教学设计",
            "检索摘要.md、course-description.md、syllabus.md、assignments.md、projects.md、instructor-insights.md",
        ),
    ]
    for row_index, values in enumerate(data, 1):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(values):
            set_cell_text(row.cells[i], value, bold=(i == 0), size=10.2)
            set_cell_width(row.cells[i], widths[i])
            if row_index % 2 == 0:
                set_cell_shading(row.cells[i], PALE_BLUE)

    doc.add_heading("共通规则", level=1)
    rules = [
        "模式决定任务工作流、默认证据范围和回答结构，知识库负责提供可引用证据。",
        "用户可在检索范围菜单中增加、取消或覆盖当前模式的默认知识库。",
        "引用内容与来源编号对应，有原文网址时允许用户直接追溯。",
        "新生成的 MIT 中文检索摘要需要重新构建 teaching_cases 向量库后才会参与课程匹配。",
    ]
    for item in rules:
        p = doc.add_paragraph(style="List Bullet")
        style_paragraph(p, space_after=4, line_spacing=1.2)
        run = p.add_run(item)
        set_run_font(run, size=11)


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    configure_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(8)
    title.add_run("北大问教四种问答模式设计说明")

    subtitle = doc.add_paragraph()
    style_paragraph(subtitle, space_after=14, line_spacing=1.2)
    run = subtitle.add_run("面向对象 工作流 知识库 返回内容与格式")
    set_run_font(run, size=11, color=MUTED)

    opening = doc.add_paragraph(
        "本文档说明北大问教中四种问答模式的用户定位和处理边界。四种模式共用同一个问答入口，"
        "但分别使用不同的检索策略、默认知识库和回答结构。通用问答处理开放式咨询；大纲设计和教学设计"
        "使用课程或案例匹配；学情诊断以教师提供的数据和评估报告为优先证据。"
    )
    style_paragraph(opening, space_after=12, line_spacing=1.3)

    doc.add_heading("模式总览", level=1)
    add_summary_table(doc)

    for mode in MODES:
        add_mode_table(doc, mode)

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_loading_table(doc)

    add_page_number(doc.sections[0].footer.paragraphs[0])

    core = doc.core_properties
    core.title = "北大问教四种问答模式设计说明"
    core.subject = "通用问答 大纲设计 教学设计 学情诊断"
    core.author = "北大问教"
    core.keywords = "智能问答, RAG, 课程大纲, 教学设计, 学情诊断"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
