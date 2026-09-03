"""评估辅助材料：课程大纲文本抽取、课件翻成页面图像。

教师选填。有大纲就对照大纲安排评本节课完整性；有课件就以课件文件评版式，
不再拿教室录像截帧当课件设计的权威证据。
"""

from __future__ import annotations

import contextlib
import shutil
import subprocess
import zipfile
from pathlib import Path

from ..config import settings
from ..rag.loaders import SUPPORTED as DOC_SUPPORTED
from ..rag.loaders import UnsupportedDocument, load_document

SYLLABUS_EXT = set(DOC_SUPPORTED)
COURSEWARE_EXT = {".pdf", ".pptx", ".ppt"}

MAX_SYLLABUS_CHARS = 12000
MAX_COURSEWARE_TEXT_CHARS = 8000

SOFFICE_CANDIDATES = (
    "soffice",
    "libreoffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice",
    "/usr/local/bin/soffice",
)


class MaterialError(RuntimeError):
    pass


def load_syllabus_text(path: str | Path, max_chars: int = MAX_SYLLABUS_CHARS) -> str:
    """把大纲文档抽成纯文本。解析失败直接抛错，避免静默丢掉教师上传的参照。"""
    path = Path(path)
    try:
        text = load_document(path)
    except UnsupportedDocument as e:
        raise MaterialError(str(e)) from e
    except FileNotFoundError as e:
        raise MaterialError(f"课程大纲文件不存在：{path.name}") from e
    text = text.strip()
    if not text:
        raise MaterialError(f"{path.name} 解析结果为空。扫描版 PDF 请先做 OCR，或改传 Word。")
    if len(text) > max_chars:
        text = text[:max_chars] + "\n……（大纲已截断）"
    return text


def _find_soffice() -> str | None:
    for cand in SOFFICE_CANDIDATES:
        p = Path(cand)
        if p.is_file():
            return str(p)
        found = shutil.which(cand)
        if found:
            return found
    return None


def _pick_indices(n: int, k: int) -> list[int]:
    if n <= 0:
        return []
    if n <= k:
        return list(range(n))
    if k == 1:
        return [n // 2]
    return [int(i * (n - 1) / (k - 1) + 0.5) for i in range(k)]


def is_valid_pptx(path: Path) -> bool:
    """扩展名是 .pptx 且能当 Office Open XML 打开（有 zip 中央目录）。"""
    if path.suffix.lower() != ".pptx" or not path.is_file():
        return False
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
        return "[Content_Types].xml" in names or any(n.startswith("ppt/") for n in names)
    except Exception:
        return False


def _office_fallbacks(path: Path) -> list[Path]:
    """同目录、同主名、不同后缀。假 pptx 旁边经常有一份能用的 .ppt。"""
    alts: list[Path] = []
    seen: set[Path] = {path.resolve()}
    for ext in (".ppt", ".pptx", ".pdf"):
        cand = path.with_suffix(ext)
        try:
            resolved = cand.resolve()
        except OSError:
            continue
        if resolved in seen or not cand.is_file():
            continue
        seen.add(resolved)
        alts.append(cand)
    return alts


def _prefer_usable_courseware(path: Path) -> tuple[Path, str]:
    """截断/损坏的 pptx 立刻改用旁边能打开的文件，避免白跑一遍 LibreOffice。"""
    if path.suffix.lower() != ".pptx" or is_valid_pptx(path):
        return path, ""
    for alt in _office_fallbacks(path):
        if alt.suffix.lower() == ".pptx" and not is_valid_pptx(alt):
            continue
        return alt, f"{path.name} 不是完整 PPTX（截断或损坏），改用 {alt.name}"
    return path, ""


def _office_to_pdf(src: Path, dest_dir: Path) -> Path:
    soffice = _find_soffice()
    if not soffice:
        raise MaterialError(
            f"无法将 {src.suffix} 转成页面图像：本机未安装 LibreOffice。"
            "请另存为 PDF 后上传，或安装 LibreOffice（brew install --cask libreoffice）。"
        )
    dest_dir.mkdir(parents=True, exist_ok=True)
    profile = dest_dir / "_lo_profile"
    profile.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        [
            soffice,
            "--headless",
            "--norestore",
            "--nolockcheck",
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(dest_dir.resolve()),
            str(src.resolve()),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=300,
    )
    expected = dest_dir / f"{src.stem}.pdf"
    if expected.is_file():
        return expected
    pdfs = sorted(p for p in dest_dir.glob("*.pdf") if p.is_file())
    if pdfs:
        return pdfs[0]
    err = (result.stderr or result.stdout or "").strip()[-400:]
    raise MaterialError(f"{src.name} 转为 PDF 失败。{err}")


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                t = (shape.text_frame.text or "").strip()
                if t:
                    texts.append(t)
        if texts:
            parts.append(f"【第 {i} 页】\n" + "\n".join(texts))
    return "\n\n".join(parts).strip()


@contextlib.contextmanager
def _mute_mupdf():
    """LibreOffice 转出的 PDF 常带 Screen 注释（嵌入媒体播放器），
    MuPDF 画不出 appearance stream，会往 stderr 刷 unsupported error。
    页面图像本来就不需要这些控件，关掉报错即可。"""
    import fitz

    tools = getattr(fitz, "TOOLS", None)
    if tools is None:
        yield
        return
    prev = tools.mupdf_display_errors()
    tools.mupdf_display_errors(False)
    try:
        tools.mupdf_display_warnings(False)
    except Exception:
        pass
    try:
        yield
    finally:
        tools.mupdf_display_errors(prev)


def _render_pdf_pages(
    pdf_path: Path, out_dir: Path, max_pages: int, prefix: str = "cw"
) -> tuple[list[dict], str]:
    import fitz

    with _mute_mupdf():
        doc = fitz.open(str(pdf_path))
        try:
            n = doc.page_count
            if n <= 0:
                raise MaterialError(f"{pdf_path.name} 没有可渲染的页面")
            indices = _pick_indices(n, max_pages)
            out_dir.mkdir(parents=True, exist_ok=True)
            pages: list[dict] = []
            text_parts: list[str] = []
            for order, page_i in enumerate(indices):
                page = doc[page_i]
                pix = page.get_pixmap(
                    matrix=fitz.Matrix(1.5, 1.5), alpha=False, annots=False
                )
                out_path = out_dir / f"{prefix}_{order:03d}.jpg"
                pix.save(str(out_path))
                pages.append(
                    {
                        "index": order,
                        "page": page_i + 1,
                        "page_count": n,
                        "path": str(out_path),
                    }
                )
                page_text = (page.get_text("text") or "").strip()
                if page_text:
                    text_parts.append(f"【第 {page_i + 1} 页】\n{page_text}")
            full_text = "\n\n".join(text_parts).strip()
            return pages, full_text
        finally:
            doc.close()


def prepare_courseware(
    path: str | Path,
    out_dir: str | Path,
    max_pages: int | None = None,
) -> dict:
    """把课件变成「页面图像 + 抽取文本」。

    PDF 直接渲染。PPT/PPTX 优先用 LibreOffice 转 PDF 再渲染；转不成时
    PPTX 退回纯文本（版式无法观测），PPT 则报错请教师另存。
    扩展名是 .pptx 但 zip 截断时，会改用同目录同名的 .ppt / .pdf。
    """
    path = Path(path)
    out_dir = Path(out_dir)
    suffix = path.suffix.lower()
    if suffix not in COURSEWARE_EXT:
        raise MaterialError(
            f"不支持的课件格式 {suffix}，目前支持：{', '.join(sorted(COURSEWARE_EXT))}"
        )
    if not path.is_file():
        raise MaterialError(f"课件文件不存在：{path.name}")

    eval_conf = settings.section("evaluation")
    max_pages = max_pages or int(eval_conf.get("courseware_max_pages", 16))
    original_name = path.name
    path, warning = _prefer_usable_courseware(path)
    suffix = path.suffix.lower()
    visual = True
    pdf_path = path

    if suffix in {".pptx", ".ppt"}:
        last_err: MaterialError | None = None
        converted = False
        sources = [path] + [
            a for a in _office_fallbacks(path) if a.suffix.lower() in {".ppt", ".pptx"}
        ]
        for src in sources:
            try:
                pdf_path = _office_to_pdf(src, out_dir / "_convert")
                if src != path:
                    note = f"{original_name} 无法转换，改用 {src.name}"
                    warning = f"{warning}；{note}" if warning else note
                    path = src
                    suffix = src.suffix.lower()
                converted = True
                break
            except MaterialError as e:
                last_err = e
        if not converted:
            if suffix == ".pptx" or path.suffix.lower() == ".pptx":
                try:
                    text = _pptx_text(path)
                except Exception as inner:
                    raise MaterialError(
                        f"{original_name} 无法转成页面图像，文本抽取也失败：{inner}"
                    ) from inner
                if not text:
                    raise MaterialError(
                        f"{original_name} 无法转成页面图像，且抽不到幻灯片文字。"
                        "请另存为 PDF 后上传。"
                    ) from last_err
                if len(text) > MAX_COURSEWARE_TEXT_CHARS:
                    text = text[:MAX_COURSEWARE_TEXT_CHARS] + "\n……（课件文本已截断）"
                return {
                    "pages": [],
                    "text": text,
                    "visual": False,
                    "page_count": 0,
                    "source": original_name,
                    "warning": str(last_err) if last_err else warning,
                }
            if last_err:
                raise last_err
            raise MaterialError(f"{original_name} 转为 PDF 失败。")

    try:
        pages, text = _render_pdf_pages(pdf_path, out_dir, max_pages)
    except MaterialError:
        raise
    except Exception as e:
        raise MaterialError(f"渲染课件页面失败：{e}") from e

    if not text and suffix == ".pptx":
        try:
            text = _pptx_text(path)
        except Exception:
            pass

    if len(text) > MAX_COURSEWARE_TEXT_CHARS:
        text = text[:MAX_COURSEWARE_TEXT_CHARS] + "\n……（课件文本已截断）"

    if not pages:
        visual = False
        extra = "课件没有可渲染的页面，仅保留抽取文本"
        warning = f"{warning}；{extra}" if warning else extra

    return {
        "pages": pages,
        "text": text,
        "visual": visual,
        "page_count": pages[0]["page_count"] if pages else 0,
        "source": path.name,
        "warning": warning,
    }
