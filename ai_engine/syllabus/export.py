"""将优化后的大纲 XML 导出为 Markdown / Word / PDF。"""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from lxml import etree

from .optimizer import detect_output_format

EXPORT_FORMATS = ("markdown", "docx", "pdf")

# 常见标签的中文展示名（未命中则 snake_case → 空格）
TAG_LABELS: dict[str, str] = {
    "course_syllabus": "课程大纲",
    "syllabus": "课程大纲",
    "course": "课程信息",
    "course_info": "课程信息",
    "name": "课程名称",
    "learning_objectives": "教学目标",
    "objective": "教学目标",
    "chapters": "教学内容",
    "chapter": "章节",
    "points": "要点",
    "point": "要点",
    "teaching_format": "教学形式",
    "summary": "说明",
    "module": "模块",
    "unit": "单元",
    "week": "周次",
    "assessment": "考核方式",
    "item": "考核项",
    "references": "参考文献",
    "reference": "文献",
    "improvements": "优化说明",
    "improvements_and_ai_strategies": "优化说明与 AI 策略",
    "ai_integration": "AI 辅助教学",
    "ai_assisted_learning": "AI 辅助学习",
    "teaching_activities": "教学活动",
    "activity": "教学活动",
    "keypoints": "教学重点",
    "difficulties": "教学难点",
    "point": "要点",
    "description": "说明",
    "change": "改动内容",
    "rationale": "依据",
}


class ExportError(RuntimeError):
    pass


def _tag_label(tag: str) -> str:
    if tag in TAG_LABELS:
        return TAG_LABELS[tag]
    return tag.replace("_", " ")


def _format_attrs(attrib: dict[str, str]) -> str:
    if not attrib:
        return ""
    hidden_metadata = {"repeatable", "origin", "source_id"}
    parts = [f"{k}={v}" for k, v in attrib.items() if k not in hidden_metadata]
    return f"（{'，'.join(parts)}）" if parts else ""


def _collect_inline_text(elem: etree._Element) -> str:
    chunks: list[str] = []
    if elem.text and elem.text.strip():
        chunks.append(elem.text.strip())
    for child in elem:
        if child.text and child.text.strip():
            chunks.append(child.text.strip())
        if child.tail and child.tail.strip():
            chunks.append(child.tail.strip())
    return "\n".join(chunks)


def _parse_xml(xml: str) -> etree._Element:
    try:
        return etree.fromstring(xml.encode("utf-8"))
    except etree.XMLSyntaxError as e:
        raise ExportError(f"XML 语法错误，无法导出：{e}") from e


def _markdown_lines(elem: etree._Element, depth: int) -> list[str]:
    lines: list[str] = []
    tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
    attrs = _format_attrs(dict(elem.attrib))
    children = [c for c in elem if isinstance(c.tag, str)]
    own_text = (elem.text or "").strip()

    if not children:
        body = own_text or _collect_inline_text(elem)
        if not body:
            return lines
        if depth == 0:
            lines.append(f"# {_tag_label(tag)}{attrs}\n")
            lines.append(body)
        else:
            level = min(depth + 1, 6)
            lines.append(f"{'#' * level} {_tag_label(tag)}{attrs}\n")
            lines.append(body)
        return lines

    if depth == 0:
        title = own_text or _tag_label(tag)
        lines.append(f"# {title}{attrs}\n")
    else:
        level = min(depth + 1, 6)
        heading = f"{'#' * level} {_tag_label(tag)}{attrs}"
        lines.append(f"{heading}\n")
        if own_text:
            lines.append(own_text)

    for child in children:
        lines.extend(_markdown_lines(child, depth + 1))
        tail = (child.tail or "").strip()
        if tail:
            lines.append(tail)
    return lines


def xml_to_markdown(xml: str) -> str:
    root = _parse_xml(xml)
    parts = _markdown_lines(root, depth=0)
    text = "\n\n".join(p.strip() for p in parts if p and p.strip())
    return text.strip() + "\n"


def _docx_add_heading(doc, text: str, level: int) -> None:
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.shared import Pt

    level = max(0, min(level, 4))
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "PingFang SC"
        run.font.size = Pt(16 - level * 2)
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT


def _docx_add_paragraph(doc, text: str) -> None:
    from docx.shared import Pt

    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = "PingFang SC"
        run.font.size = Pt(11)


def _docx_walk(doc, elem: etree._Element, depth: int) -> None:
    tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
    attrs = _format_attrs(dict(elem.attrib))
    children = [c for c in elem if isinstance(c.tag, str)]
    own_text = (elem.text or "").strip()

    if not children:
        body = own_text or _collect_inline_text(elem)
        if not body:
            return
        if depth == 0:
            _docx_add_heading(doc, f"{_tag_label(tag)}{attrs}", 0)
        else:
            _docx_add_heading(doc, f"{_tag_label(tag)}{attrs}", min(depth, 4))
        for para in body.split("\n"):
            if para.strip():
                _docx_add_paragraph(doc, para.strip())
        return

    if depth == 0:
        title = own_text or _tag_label(tag)
        _docx_add_heading(doc, f"{title}{attrs}", 0)
    else:
        _docx_add_heading(doc, f"{_tag_label(tag)}{attrs}", min(depth, 4))
        if own_text:
            _docx_add_paragraph(doc, own_text)

    for child in children:
        _docx_walk(doc, child, depth + 1)
        tail = (child.tail or "").strip()
        if tail:
            _docx_add_paragraph(doc, tail)


def xml_to_docx(xml: str) -> bytes:
    from docx import Document

    doc = Document()
    _docx_walk(doc, _parse_xml(xml), depth=0)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _run_cmd(cmd: list[str], *, cwd: Path | None = None) -> None:
    try:
        out = subprocess.run(
            cmd,
            capture_output=True,
            timeout=120,
            cwd=str(cwd) if cwd else None,
        )
    except FileNotFoundError as e:
        raise ExportError(f"未找到命令：{cmd[0]}") from e
    except subprocess.TimeoutExpired as e:
        raise ExportError(f"命令超时：{cmd[0]}") from e
    if out.returncode != 0:
        err = (out.stderr or out.stdout or b"").decode("utf-8", errors="replace")
        raise ExportError(err.strip() or f"命令失败：{cmd[0]}")


def _find_soffice() -> str | None:
    candidates = [
        "soffice",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]
    for path in candidates:
        if shutil.which(path) or Path(path).exists():
            return path
    return None


def _markdown_to_pdf_pandoc(md_path: Path, pdf_path: Path) -> None:
    if not shutil.which("pandoc"):
        raise ExportError("pandoc 未安装")
    engines = [
        ["xelatex", "PingFang SC"],
        ["xelatex", "Songti SC"],
        ["pdflatex", None],
    ]
    last_err: Exception | None = None
    for engine, font in engines:
        cmd = [
            "pandoc",
            str(md_path),
            "-o",
            str(pdf_path),
            f"--pdf-engine={engine}",
        ]
        if font:
            cmd.extend(["-V", f"CJKmainfont={font}", "-V", "geometry:margin=2.5cm"])
        try:
            _run_cmd(cmd)
            if pdf_path.exists() and pdf_path.stat().st_size > 0:
                return
        except ExportError as e:
            last_err = e
    if last_err:
        raise last_err
    raise ExportError("pandoc 未能生成 PDF")


def _docx_to_pdf_soffice(docx_path: Path, out_dir: Path) -> Path:
    soffice = _find_soffice()
    if not soffice:
        raise ExportError("LibreOffice (soffice) 未安装")
    _run_cmd([soffice, "--headless", "--convert-to", "pdf", str(docx_path)], cwd=out_dir)
    pdf_path = out_dir / f"{docx_path.stem}.pdf"
    if not pdf_path.exists():
        raise ExportError("LibreOffice 未能生成 PDF")
    return pdf_path


def xml_to_pdf(xml: str) -> bytes:
    md = xml_to_markdown(xml)
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        md_path = root / "syllabus.md"
        pdf_path = root / "syllabus.pdf"
        docx_path = root / "syllabus.docx"
        md_path.write_text(md, encoding="utf-8")

        if shutil.which("pandoc"):
            try:
                _markdown_to_pdf_pandoc(md_path, pdf_path)
                return pdf_path.read_bytes()
            except ExportError:
                pass

        docx_path.write_bytes(xml_to_docx(xml))
        try:
            pdf_path = _docx_to_pdf_soffice(docx_path, root)
            return pdf_path.read_bytes()
        except ExportError as e:
            raise ExportError(
                "PDF 导出需要安装 pandoc 或 LibreOffice。"
                "macOS 可执行：brew install pandoc 或 brew install --cask libreoffice。"
                f"（{e}）"
            ) from e


def markdown_to_docx(markdown: str) -> bytes:
    if not shutil.which("pandoc"):
        raise ExportError("docx 导出需要安装 pandoc（brew install pandoc）")
    with tempfile.TemporaryDirectory() as tmp:
        md_path = Path(tmp) / "syllabus.md"
        docx_path = Path(tmp) / "syllabus.docx"
        md_path.write_text(markdown.strip() + "\n", encoding="utf-8")
        _run_cmd(["pandoc", str(md_path), "-o", str(docx_path)])
        return docx_path.read_bytes()


def markdown_to_pdf(markdown: str) -> bytes:
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        md_path = root / "syllabus.md"
        pdf_path = root / "syllabus.pdf"
        docx_path = root / "syllabus.docx"
        md_path.write_text(markdown.strip() + "\n", encoding="utf-8")

        if shutil.which("pandoc"):
            try:
                _markdown_to_pdf_pandoc(md_path, pdf_path)
                return pdf_path.read_bytes()
            except ExportError:
                pass

        docx_path.write_bytes(markdown_to_docx(markdown))
        return _docx_to_pdf_soffice(docx_path, root).read_bytes()


def export_syllabus(content: str, fmt: str) -> tuple[bytes, str, str]:
    """导出为 (bytes, media_type, filename_suffix)。"""
    fmt = fmt.lower()
    if fmt not in EXPORT_FORMATS:
        raise ExportError(f"不支持的格式 {fmt}，可选：{', '.join(EXPORT_FORMATS)}")

    if detect_output_format(content) == "markdown":
        if fmt == "markdown":
            return (
                content.strip().encode("utf-8") + b"\n",
                "text/markdown; charset=utf-8",
                ".md",
            )
        if fmt == "docx":
            return (
                markdown_to_docx(content),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ".docx",
            )
        return markdown_to_pdf(content), "application/pdf", ".pdf"

    if fmt == "markdown":
        content_bytes = xml_to_markdown(content).encode("utf-8")
        return content_bytes, "text/markdown; charset=utf-8", ".md"
    if fmt == "docx":
        return (
            xml_to_docx(content),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".docx",
        )
    return xml_to_pdf(content), "application/pdf", ".pdf"


def content_disposition(filename: str) -> str:
    safe = re.sub(r'[\\/:*?"<>|]', "_", filename) or "syllabus"
    ascii_fallback = safe.encode("ascii", "ignore").decode() or "syllabus"
    return (
        f'attachment; filename="{ascii_fallback}"; '
        f"filename*=UTF-8''{quote(safe)}"
    )


def export_filename(course_name: str, suffix: str) -> str:
    base = (course_name or "syllabus").replace("/", "_").strip()
    return f"{base}{suffix}"
