"""文档加载：把 pdf / docx / md / txt / html 统一读成纯文本。"""

from __future__ import annotations

import re
import subprocess
from pathlib import Path
from urllib.parse import urlparse

SUPPORTED = {".pdf", ".docx", ".md", ".txt", ".html", ".htm", ".doc"}

# 不入库的目录/文件名（原始视频、转写中间产物、说明文档、演示种子）
SKIP_DIR_NAMES = {
    "raw_video",
    ".work",
    "__pycache__",
    # 原始网页/PDF 快照只用于追溯，避免与清洗后的正文重复入库。
    "source",
    "original_sources",
}
SKIP_FILE_NAMES = {"readme.md"}
SKIP_FILE_PREFIXES = ("种子-",)


class UnsupportedDocument(ValueError):
    pass


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


_CJK_RE = re.compile(r"[\u4e00-\u9fff\u3400-\u4dbf]")
# pypdf 对部分中文 PDF 的 CMap 映射错误时，常出现 Malayalam/Bengali 等无关脚本字符
_EXOTIC_SCRIPT_RE = re.compile(r"[\u0980-\u0dff\u10a0-\u10ff\u1200-\u137f]")


def _cjk_count(text: str) -> int:
    return len(_CJK_RE.findall(text))


def _exotic_script_count(text: str) -> int:
    return len(_EXOTIC_SCRIPT_RE.findall(text))


def _pdf_text_quality(text: str) -> float:
    cjk = _cjk_count(text)
    exotic = _exotic_script_count(text)
    return cjk * 10 - exotic * 8 + min(len(text), 5000) * 0.001


def _pdf_looks_garbled(text: str) -> bool:
    if len(text) < 100:
        return False
    cjk = _cjk_count(text)
    exotic = _exotic_script_count(text)
    return cjk < 3 and exotic > 10


def _load_pdf_pymupdf(path: Path) -> str:
    import fitz

    doc = fitz.open(str(path))
    try:
        return "\n\n".join(page.get_text("text") for page in doc)
    finally:
        doc.close()


def _load_pdf_pypdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _load_pdf_pdftotext(path: Path) -> str:
    out = subprocess.run(
        ["pdftotext", "-enc", "UTF-8", "-layout", str(path), "-"],
        capture_output=True,
        timeout=120,
    )
    if out.returncode != 0 or not out.stdout:
        raise RuntimeError("pdftotext failed")
    return out.stdout.decode("utf-8", errors="replace")


def _load_pdf(path: Path) -> str:
    """多后端抽取 PDF 文本，按中文质量评分选最优结果。

    pypdf 对部分 CID/CMap 中文 PDF 会把汉字映射成无关 Unicode 字符；
    PyMuPDF 通常能正确读取同一文件的文本层。
    """
    backends = [
        ("pymupdf", _load_pdf_pymupdf),
        ("pdftotext", _load_pdf_pdftotext),
        ("pypdf", _load_pdf_pypdf),
    ]

    candidates: list[tuple[str, str]] = []
    for name, loader in backends:
        try:
            text = _clean(loader(path))
            if text:
                candidates.append((name, text))
        except (FileNotFoundError, ImportError, RuntimeError, OSError):
            continue

    if not candidates:
        return ""

    _best_name, best_text = max(candidates, key=lambda item: _pdf_text_quality(item[1]))

    if _pdf_looks_garbled(best_text):
        tried = ", ".join(name for name, _ in candidates)
        raise UnsupportedDocument(
            f"{path.name} 的中文 PDF 解析失败（已尝试 {tried}）。"
            "请改传 .docx，或使用「粘贴文本」。扫描版 PDF 需先做 OCR。"
        )

    return best_text


def _load_docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # 大纲类文档的学时/考核权重常放在表格里，漏掉会丢关键信息
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _load_doc(path: Path) -> str:
    """老式 .doc 二进制格式，尝试用系统 textutil(macOS) 转换。"""
    try:
        out = subprocess.run(
            ["textutil", "-convert", "txt", "-stdout", str(path)],
            capture_output=True,
            timeout=60,
        )
        if out.returncode == 0 and out.stdout:
            return out.stdout.decode("utf-8", errors="ignore")
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass
    raise UnsupportedDocument(
        f"无法解析旧版 .doc 文件 {path.name}。请在 Word 中另存为 .docx 或 PDF 后重试。"
    )


def _load_html(path: Path) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    return soup.get_text("\n")


_SOURCE_URL_RE = re.compile(
    r"(?:原始\s*URL|Source\s*URL|原文链接|来源网址)\s*[:：]\s*"
    r"<?(https?://[^\s)>\]|。，；;]+)",
    re.IGNORECASE,
)
_MANIFEST_ROW_RE = re.compile(
    r"\|\s*`([^`]+)`\s*\|\s*<?(https?://[^\s>|]+)", re.IGNORECASE
)
_LOCAL_FILE_ROW_RE = re.compile(r"\|\s*\*\*本地文件\*\*\s*\|\s*`([^`]+)`")
_ORIGINAL_URL_ROW_RE = re.compile(
    r"\|\s*\*\*原始\s*URL\*\*\s*\|\s*<?(https?://[^\s>|]+)", re.IGNORECASE
)


def _valid_web_url(value: str | None) -> str:
    """只接受可供用户跳转的 HTTP(S) 来源，避免将页内脚本地址误当来源。"""
    value = (value or "").strip().rstrip(".,;:，。；：")
    parsed = urlparse(value)
    if parsed.scheme in {"http", "https"} and parsed.netloc:
        return value
    return ""


def _html_source_url(path: Path) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    canonical = soup.find("link", rel=lambda value: value and "canonical" in value)
    if canonical:
        url = _valid_web_url(canonical.get("href"))
        if url:
            return url
    og_url = soup.find("meta", attrs={"property": "og:url"})
    if og_url:
        return _valid_web_url(og_url.get("content"))
    return ""


def _manifest_source_urls(root: Path) -> dict[Path, str]:
    """读取语料目录 README 中已有的「本地文件—原始 URL」对照表。"""
    mapping: dict[Path, str] = {}
    for readme in root.rglob("README.md"):
        text = readme.read_text(encoding="utf-8", errors="ignore")

        # 归档清单：| `file.html` | https://... |
        for match in _MANIFEST_ROW_RE.finditer(text):
            mapping[(readme.parent / match.group(1)).resolve()] = _valid_web_url(
                match.group(2)
            )

        # 课程大纲目录：「本地文件」和「原始 URL」分属两行。
        pending_file: str | None = None
        for line in text.splitlines():
            file_match = _LOCAL_FILE_ROW_RE.search(line)
            if file_match:
                pending_file = file_match.group(1)
                continue
            url_match = _ORIGINAL_URL_ROW_RE.search(line)
            if pending_file and url_match:
                mapping[(readme.parent / pending_file).resolve()] = _valid_web_url(
                    url_match.group(1)
                )
                pending_file = None
    return mapping


def source_url_for_document(path: str | Path, root: str | Path, text: str = "") -> str:
    """找出文档的原始公网地址；无公网来源时返回空字符串。"""
    path = Path(path)
    root = Path(root)

    if path.suffix.lower() in {".html", ".htm"}:
        url = _html_source_url(path)
        if url:
            return url

    marked = _SOURCE_URL_RE.search(text)
    if marked:
        url = _valid_web_url(marked.group(1))
        if url:
            return url

    manifest = _manifest_source_urls(root)
    direct = manifest.get(path.resolve(), "")
    if direct:
        return direct

    # 整理稿常在正文中指向 source/ 下的原件，继承原件的 URL。
    for local_ref in re.findall(r"`([^`]+\.(?:html?|pdf))`", text, re.IGNORECASE):
        inherited = manifest.get((path.parent / local_ref).resolve(), "")
        if inherited:
            return inherited
    return ""


def load_document(path: str | Path) -> str:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在: {path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _load_pdf(path)
    elif suffix == ".docx":
        text = _load_docx(path)
    elif suffix == ".doc":
        text = _load_doc(path)
    elif suffix in (".html", ".htm"):
        text = _load_html(path)
    elif suffix in (".md", ".txt"):
        text = path.read_text(encoding="utf-8", errors="ignore")
    else:
        raise UnsupportedDocument(
            f"不支持的文件格式 {suffix}，目前支持：{', '.join(sorted(SUPPORTED))}"
        )

    text = _clean(text)
    if not text:
        raise UnsupportedDocument(
            f"{path.name} 解析结果为空。若是扫描版 PDF，需要先做 OCR。"
        )
    return text


def _should_skip(path: Path, root: Path) -> bool:
    rel = path.relative_to(root)
    for part in rel.parts[:-1]:
        if part in SKIP_DIR_NAMES or part.startswith("."):
            return True
    name_lower = path.name.lower()
    if name_lower in SKIP_FILE_NAMES:
        return True
    if any(path.name.startswith(p) for p in SKIP_FILE_PREFIXES):
        return True
    return False


def iter_documents(directory: str | Path):
    """递归遍历目录下所有支持的文档，产出 (path, text)。"""
    directory = Path(directory)
    if not directory.exists():
        return
    for path in sorted(directory.rglob("*")):
        if not path.is_file() or path.name.startswith("."):
            continue
        if _should_skip(path, directory):
            continue
        if path.suffix.lower() not in SUPPORTED:
            continue
        try:
            yield path, load_document(path)
        except Exception as e:  # 单个文件坏了不该中断整个建库
            print(f"  [跳过] {path.name}: {e}")
